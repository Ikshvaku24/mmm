"""A benchmark-comparison sheet you fill in yourself, with the formulas already in it.

You paste the vendor's (or last year's, or anyone's) contribution into one
column and every diagnostic recalculates live in the spreadsheet - the gap, the
ratio, how far the data wants to move away from your prior, and the corrected
`global_prior_mean` to write back into the prior file.

Why this rather than a Python script that reads a comparison file: the benchmark
never arrives in a fixed schema. It is a column in a deck, a tab in someone's
workbook, a number read off a slide. Handing you a sheet with the run's own
numbers already laid out and the formulas pre-written means the only step left
is paste - and you can see every intermediate quantity instead of trusting a
script's join.

The arithmetic in the formulas
------------------------------
For a sign-constrained feature the coefficient is `beta = exp(Normal(mu, sigma))`
and the reported contribution is the posterior MEDIAN, so

    contribution = exp(mu_post) * SUM(x_scaled + shift) * dv_scale

With a conjugate update `mu_post = (1-c)*mu_prior + c*mu_like`, where c is the
`contraction`. If the prior was derived FROM the benchmark, then

    ln(ours / theirs) = -c * delta

so `delta` - how far below (or above) your prior the data wants the coefficient,
in log units - is recoverable as `ln(ratio) / c`, and the prior mean that would
land on the benchmark is

    new_prior_mean = old_prior_mean * ratio^(1 / (1 - c))

**Reading `delta` across features is the point.** If it clusters around one
value, a single global constraint is pushing every driver the same way - the
usual cause being a term the benchmark does not have (a free region intercept)
claiming a share of a fixed total. Correcting prior means one at a time cannot
fix that: the next refit re-imposes the same shortfall on the corrected numbers.
If `delta` scatters, the gaps really are per-feature and the corrected means in
column K are the fix.

Output
------
`05_contributions/benchmark_comparison.xlsx` when openpyxl is available (it is
on Databricks), otherwise `benchmark_comparison.csv`. Excel evaluates `=`
formulas in a CSV on open, so the fallback behaves the same way - it just cannot
carry column widths or a second sheet.
"""
from __future__ import annotations

import os

import numpy as np
import pandas as pd

# Column layout. Order is deliberate: what you read, then what you FILL, then
# what recalculates, then the raw inputs the formulas use.
COLUMNS = [
    ("region", "the region this row belongs to"),
    ("feature", "feature name, exactly as in the prior file"),
    ("pillar", "reporting group"),
    ("our_contribution", "what this run reports, KPI units"),
    ("benchmark_contribution", "<<< PASTE THE BENCHMARK HERE >>>"),
    ("pct_diff", "(ours - theirs) / |theirs| * 100"),
    ("ratio_needed", "theirs / ours - the multiplier to close the gap"),
    ("contraction", "1 - posterior_var/prior_var. Near 0 = the posterior IS your prior"),
    ("delta", "ln(ratio)/contraction - how far the data wants to move, log units"),
    ("current_prior_mean", "the implied median coefficient your prior file produces"),
    ("suggested_prior_mean", "current * ratio^(1/(1-contraction)) - write this back"),
    ("implied_benchmark_beta", "their contribution expressed on YOUR axis"),
    ("effective_scaled_sum", "SUM(x_scaled + reference shift) over the window"),
    ("dv_scale_used", "the KPI scale for this region"),
    ("verdict", "what to do about this row"),
]
FILL_COL = "benchmark_contribution"


def _letters():
    return {name: chr(ord("A") + i) for i, (name, _) in enumerate(COLUMNS)}


def formulas_for_row(r: int) -> dict:
    """The formula for every computed column, for spreadsheet row `r` (1-based).

    Every one is guarded on the benchmark cell being empty, so an unfilled sheet
    shows blanks rather than #DIV/0! everywhere.
    """
    L = _letters()
    ours, bench = f"{L['our_contribution']}{r}", f"{L[FILL_COL]}{r}"
    ratio, contr = f"{L['ratio_needed']}{r}", f"{L['contraction']}{r}"
    blank = f'{bench}="",'
    return {
        "pct_diff": f'=IF(OR({blank}{bench}=0),"",({ours}-{bench})/ABS({bench})*100)',
        "ratio_needed": f'=IF(OR({blank}{ours}=0),"",{bench}/{ours})',
        # delta needs a genuine contraction in (0,1): at c<=0 the posterior is
        # wider than the prior (unidentified) and the back-out is meaningless.
        "delta": (f'=IF(OR({blank}{ratio}<=0,{contr}<=0.001,{contr}>=0.999),"",'
                  f'LN({ratio})/{contr})'),
        "suggested_prior_mean": (
            f'=IF(OR({blank}{ratio}<=0,{contr}>=0.999),"",'
            f'{L["current_prior_mean"]}{r}*{ratio}^(1/(1-{contr})))'),
        "implied_benchmark_beta": (
            f'=IF(OR({blank}{L["effective_scaled_sum"]}{r}=0,'
            f'{L["dv_scale_used"]}{r}=0),"",'
            f'{bench}/({L["effective_scaled_sum"]}{r}*{L["dv_scale_used"]}{r}))'),
        "verdict": (
            f'=IF({bench}="","paste the benchmark",'
            f'IF(ABS({L["pct_diff"]}{r})<10,"ok",'
            f'IF({contr}<=0,"UNIDENTIFIED - structural, no prior fixes it",'
            f'IF({contr}<0.2,"prior-driven - write column K back to the prior file",'
            f'IF({contr}>0.5,"data disagrees - impose (tighten sd) or accept the gap",'
            f'"mixed - check delta against the other rows")))))'),
    }
COMPUTED = set(formulas_for_row(2))


def summary_formulas(first: int, last: int) -> list:
    """(label, formula) pairs for the read-first block beside the table."""
    L = _letters()
    pd_rng = f"{L['pct_diff']}{first}:{L['pct_diff']}{last}"
    dl_rng = f"{L['delta']}{first}:{L['delta']}{last}"
    return [
        ("rows filled in", f'=COUNT({L[FILL_COL]}{first}:{L[FILL_COL]}{last})'),
        ("median % diff", f'=IFERROR(MEDIAN({pd_rng}),"")'),
        ("features BELOW the benchmark", f'=COUNTIF({pd_rng},"<0")'),
        ("features ABOVE the benchmark", f'=COUNTIF({pd_rng},">0")'),
        ("median delta", f'=IFERROR(MEDIAN({dl_rng}),"")'),
        ("delta spread (IQR)",
         f'=IFERROR(QUARTILE({dl_rng},3)-QUARTILE({dl_rng},1),"")'),
        ("implied shrink vs your priors",
         f'=IFERROR(TEXT(EXP(-MEDIAN({dl_rng}))-1,"0.0%"),"")'),
        ("reading", '=IF(COUNT(' + f'{L["delta"]}{first}:{L["delta"]}{last}'
         + ')<3,"need >=3 filled rows",'
         'IF(IFERROR(QUARTILE(' + dl_rng + ',3)-QUARTILE(' + dl_rng + ',1),9)<0.1,'
         '"delta CLUSTERS -> ONE global constraint (usually a free intercept the '
         'benchmark does not have). Fixing prior means one by one will not hold.",'
         '"delta SCATTERS -> per-feature prior errors. Column K is the fix."))'),
    ]


# --------------------------------------------------------------------------- #
# assembling the run's own numbers
# --------------------------------------------------------------------------- #
def build_table(decomp, pdata, outdir_root: str = None,
                math_df: pd.DataFrame = None,
                contraction_df: pd.DataFrame = None,
                prior_df: pd.DataFrame = None) -> pd.DataFrame:
    """One row per region x feature, with every input column filled and the
    computed columns left empty (the formulas go in at write time)."""
    if math_df is None:
        raise ValueError("build_table needs contribution_math (math_df)")
    m = math_df.copy()
    ren = {"volume_median_of_total": "our_contribution",
           "dv_scale_used": "dv_scale_used"}
    m = m.rename(columns={k: v for k, v in ren.items() if k in m.columns})
    keep = ["region", "feature", "pillar", "our_contribution",
            "effective_scaled_sum", "dv_scale_used"]
    out = m[[c for c in keep if c in m.columns]].copy()
    for c in keep:
        if c not in out.columns:
            out[c] = np.nan

    out["contraction"] = np.nan
    if contraction_df is not None and len(contraction_df):
        c = contraction_df
        if "role" in c.columns:
            c = c[c["role"].astype(str).str.startswith("coefficient")]
        if {"name", "contraction"} <= set(c.columns):
            c = c.assign(feature=c["name"].astype(str).str.split(" @ ").str[0]
                         .str.strip())
            agg = c.groupby("feature", as_index=False)["contraction"].mean()
            out = out.drop(columns=["contraction"]).merge(agg, on="feature",
                                                          how="left")

    out["current_prior_mean"] = np.nan
    if prior_df is not None and len(prior_df):
        p = prior_df
        fcol = "variable" if "variable" in p.columns else "feature"
        if "region" in p.columns:
            pop = p[p["region"].astype(str) == "__population__"]
            p = pop if len(pop) else p
        if fcol in p.columns and "implied_median" in p.columns:
            agg = (p.groupby(fcol, as_index=False)["implied_median"].mean()
                   .rename(columns={fcol: "feature",
                                    "implied_median": "current_prior_mean"}))
            out = out.drop(columns=["current_prior_mean"]).merge(
                agg, on="feature", how="left")

    out[FILL_COL] = np.nan
    for name in COMPUTED:
        out[name] = ""
    return out[[c for c, _ in COLUMNS]].sort_values(["region", "feature"])


# --------------------------------------------------------------------------- #
# writing
# --------------------------------------------------------------------------- #
NOTE = (
    "Paste the benchmark contribution into column E. Everything else "
    "recalculates. Read the summary block to the right FIRST: if delta "
    "clusters, the gap is one global constraint (a term the benchmark does not "
    "have) and correcting prior means one at a time will not hold. If delta "
    "scatters, column K is the corrected global_prior_mean to write back into "
    "the prior file."
)


def write_benchmark_sheet(table: pd.DataFrame, outdir: str) -> str:
    """Write the sheet with live formulas. .xlsx if openpyxl is here, else .csv."""
    os.makedirs(outdir, exist_ok=True)
    try:
        import openpyxl  # noqa: F401
        return _write_xlsx(table, outdir)
    except ImportError:
        return _write_formula_csv(table, outdir)


def _rows_with_formulas(table: pd.DataFrame):
    """Yield each data row as a list of cell values, formulas substituted in."""
    names = [c for c, _ in COLUMNS]
    for i, (_, row) in enumerate(table.iterrows()):
        r = i + 2                      # +1 for the header, +1 for 1-based
        f = formulas_for_row(r)
        cells = []
        for name in names:
            if name in f:
                cells.append(f[name])
            else:
                v = row[name]
                cells.append("" if (v is None or (isinstance(v, float)
                                                  and not np.isfinite(v))) else v)
        yield cells


def _write_formula_csv(table: pd.DataFrame, outdir: str) -> str:
    import csv

    path = os.path.join(outdir, "benchmark_comparison.csv")
    names = [c for c, _ in COLUMNS]
    first, last = 2, len(table) + 1
    with open(path, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(names + ["", "SUMMARY - read this first", ""])
        summ = summary_formulas(first, last)
        for i, cells in enumerate(_rows_with_formulas(table)):
            extra = ["", summ[i][0], summ[i][1]] if i < len(summ) else []
            w.writerow(cells + extra)
        # if there are fewer data rows than summary lines, keep the rest
        for j in range(len(table), len(summ)):
            w.writerow([""] * len(names) + ["", summ[j][0], summ[j][1]])
        w.writerow([])
        w.writerow(["NOTE", NOTE])
    return path


def _write_xlsx(table: pd.DataFrame, outdir: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path = os.path.join(outdir, "benchmark_comparison.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "comparison"
    names = [c for c, _ in COLUMNS]
    ws.append(names)
    fill_idx = names.index(FILL_COL) + 1
    hdr = Font(bold=True)
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for j, _ in enumerate(names, start=1):
        ws.cell(row=1, column=j).font = hdr
    ws.cell(row=1, column=fill_idx).fill = yellow

    for cells in _rows_with_formulas(table):
        ws.append(cells)
    for i in range(2, len(table) + 2):
        ws.cell(row=i, column=fill_idx).fill = yellow

    # summary block, two columns to the right of the table
    lab_c, val_c = len(names) + 2, len(names) + 3
    ws.cell(row=1, column=lab_c, value="SUMMARY - read this first").font = hdr
    for i, (label, formula) in enumerate(summary_formulas(2, len(table) + 1),
                                         start=2):
        ws.cell(row=i, column=lab_c, value=label)
        ws.cell(row=i, column=val_c, value=formula)
    note_row = len(summary_formulas(2, 2)) + 3
    c = ws.cell(row=note_row, column=lab_c, value=NOTE)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=note_row, start_column=lab_c,
                   end_row=note_row + 5, end_column=lab_c + 3)

    widths = {"region": 22, "feature": 46, "pillar": 16, "verdict": 52}
    for j, name in enumerate(names, start=1):
        ws.column_dimensions[get_column_letter(j)].width = widths.get(name, 17)
    ws.column_dimensions[get_column_letter(lab_c)].width = 32
    ws.column_dimensions[get_column_letter(val_c)].width = 18
    ws.freeze_panes = "C2"

    # a second sheet documenting every column, so the sheet explains itself
    doc = wb.create_sheet("what each column means")
    doc.append(["column", "meaning"])
    doc.cell(row=1, column=1).font = hdr
    doc.cell(row=1, column=2).font = hdr
    for i, (name, meaning) in enumerate(COLUMNS, start=2):
        doc.cell(row=i, column=1, value=name)
        doc.cell(row=i, column=2, value=meaning)
    doc.column_dimensions["A"].width = 26
    doc.column_dimensions["B"].width = 96
    wb.save(path)
    return path


def write_benchmark_comparison(decomp, pdata, outdir: str,
                               math_df=None, run_root: str = None) -> str:
    """Called by the pipeline: assemble and write the sheet."""
    contraction_df = prior_df = None
    if run_root:
        cp = os.path.join(run_root, "02_convergence",
                          "prior_posterior_contraction.csv")
        pp = os.path.join(run_root, "01_data", "prior_summary.csv")
        if os.path.exists(cp):
            contraction_df = pd.read_csv(cp)
        if os.path.exists(pp):
            prior_df = pd.read_csv(pp)
    table = build_table(decomp, pdata, math_df=math_df,
                        contraction_df=contraction_df, prior_df=prior_df)
    path = write_benchmark_sheet(table, outdir)
    print(f"[benchmark] paste the benchmark into column "
          f"{_letters()[FILL_COL]} of {os.path.basename(path)} "
          f"({len(table)} rows)")
    return path
